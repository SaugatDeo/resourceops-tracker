"""generate_data.py - builds 100% synthetic data for the ResourceOps Tracker.
Run from the ResourceOps folder:  python scripts/generate_data.py
Output: data/ResourceOps_Tracker_filled.xlsx and data/planted_errors_DO_NOT_OPEN.csv
"""
import random
from datetime import date, timedelta
from pathlib import Path

import pandas as pd
from faker import Faker
from openpyxl.styles import Font, PatternFill

random.seed(42)
Faker.seed(42)
fake = Faker("en_IN")

TODAY = date(2026, 10, 7)
REGIONS = ["North", "South", "East", "West"]
GRADES = ["Analyst", "Consultant", "Senior Consultant", "Manager"]
CITIES = {
    "North": ["Delhi", "Gurugram", "Lucknow"],
    "South": ["Chennai", "Bengaluru", "Hyderabad"],
    "East": ["Kolkata", "Bhubaneswar", "Patna"],
    "West": ["Mumbai", "Pune", "Ahmedabad"],
}
CLIENT_WORDS = ["Alpha", "Bravo", "Charlie", "Delta", "Echo", "Foxtrot", "Golf",
                "Hotel", "India", "Juliet", "Kilo", "Lima", "Mike", "November",
                "Oscar", "Papa", "Quebec", "Romeo", "Sierra", "Tango", "Uniform",
                "Victor", "Whiskey", "Xray", "Yankee"]
COURSES = ["Data Protection Basics", "Code of Conduct",
           "Information Security Awareness", "Anti-Bribery Essentials"]

# ---------- 1. Practitioners (40) ----------
practitioners, used = [], set()
for i in range(1, 41):
    name = fake.name()
    while name in used:
        name = fake.name()
    used.add(name)
    practitioners.append({
        "PractID": f"P{i:03d}",
        "Name": name,
        "Region": REGIONS[(i - 1) % 4],
        "Grade": random.choice(GRADES),
        "WeekendAvailable": "Yes" if random.random() < 0.7 else "No",
    })

# ---------- 2. Leave (18 clean rows) ----------
leave = []
for i in range(1, 19):
    start = date(2026, 9, 1) + timedelta(days=random.randint(0, 105))
    leave.append({
        "LeaveID": f"L{i:03d}",
        "PractID": random.choice(practitioners)["PractID"],
        "LeaveStart": start,
        "LeaveEnd": start + timedelta(days=random.randint(0, 4)),
        "LeaveType": random.choice(["Annual", "Sick", "Exam"]),
    })

def on_leave(pid, d):
    return any(l["PractID"] == pid and l["LeaveStart"] <= d <= l["LeaveEnd"] for l in leave)

# ---------- 3. Requests (25) - stock counts happen on Saturdays ----------
saturdays = [date(2026, 9, 5) + timedelta(days=7 * k) for k in range(14)]
requests = []
for i in range(1, 26):
    ev = random.choice(saturdays)
    region = random.choice(REGIONS)
    received = min(ev - timedelta(days=random.randint(10, 25)),
                   TODAY - timedelta(days=random.randint(0, 5)))
    if ev < TODAY:
        status, closed = "Closed", received + timedelta(days=random.randint(2, 9))
    else:
        status = random.choices(["Scheduled", "Open"], weights=[75, 25])[0]
        closed = None
    requests.append({
        "RequestID": f"R{i:03d}",
        "Client": f"Client {CLIENT_WORDS[i - 1]} Ltd",
        "Region": region,
        "Location": random.choice(CITIES[region]),
        "EventDate": ev,
        "PeopleNeeded": random.randint(2, 5),
        "DateReceived": received,
        "Status": status,
        "DateClosed": closed,
    })

# ---------- 4. Schedule (clean first) ----------
booked, schedule = set(), []
for r in requests:
    if r["Status"] == "Open":
        continue
    same = [p for p in practitioners if p["Region"] == r["Region"]]
    other = [p for p in practitioners if p["Region"] != r["Region"]]
    random.shuffle(same); random.shuffle(other)
    chosen = []
    for p in same + other:
        if len(chosen) == r["PeopleNeeded"]:
            break
        if p["WeekendAvailable"] != "Yes":
            continue
        if (p["PractID"], r["EventDate"]) in booked or on_leave(p["PractID"], r["EventDate"]):
            continue
        chosen.append(p)
    for p in chosen:
        booked.add((p["PractID"], r["EventDate"]))
        schedule.append({
            "ScheduleID": f"S{len(schedule) + 1:03d}",
            "RequestID": r["RequestID"],
            "PractID": p["PractID"],
            "Date": r["EventDate"],
            "Hours": random.choice([8, 8, 10]),
        })

# ---------- 5. Plant known errors (so we can test our checks later) ----------
planted, touched = [], set()
status_of = {r["RequestID"]: r["Status"] for r in requests}
date_of = {r["RequestID"]: r["EventDate"] for r in requests}

# 5a. up to 3 double-bookings: same person on two requests the same Saturday
by_date = {}
for r in requests:
    if r["Status"] != "Open":
        by_date.setdefault(r["EventDate"], []).append(r["RequestID"])
n = 0
for d, rids in by_date.items():
    if n == 3 or len(rids) < 2:
        continue
    r1, r2 = rids[0], rids[1]
    p1 = {s["PractID"] for s in schedule if s["RequestID"] == r1}
    rows2 = [s for s in schedule if s["RequestID"] == r2]
    p2 = {s["PractID"] for s in rows2}
    options = sorted(p1 - p2)
    if options and rows2:
        x, row = random.choice(options), random.choice(rows2)
        row["PractID"] = x
        touched.add(row["ScheduleID"])
        planted.append(("DOUBLE_BOOKING", f"{x} on {d:%d-%b-%Y} in {r1} and {r2}"))
        n += 1

# 5b. 2 leave clashes: person scheduled during approved leave
future_rows = [s for s in schedule
               if status_of[s["RequestID"]] == "Scheduled" and s["ScheduleID"] not in touched]
for s in random.sample(future_rows, 2):
    leave.append({
        "LeaveID": f"L{len(leave) + 1:03d}",
        "PractID": s["PractID"],
        "LeaveStart": s["Date"] - timedelta(days=1),
        "LeaveEnd": s["Date"] + timedelta(days=1),
        "LeaveType": "Annual",
    })
    touched.add(s["ScheduleID"])
    planted.append(("LEAVE_CLASH", f"{s['PractID']} scheduled {s['Date']:%d-%b-%Y} in {s['RequestID']} during leave"))

# 5c. 2 understaffed requests: remove one person
cands = sorted({s["RequestID"] for s in schedule
                if status_of[s["RequestID"]] == "Scheduled"
                and sum(1 for t in schedule if t["RequestID"] == s["RequestID"]) >= 3})
for rid in random.sample(cands, min(2, len(cands))):
    rows = [s for s in schedule if s["RequestID"] == rid and s["ScheduleID"] not in touched]
    if rows:
        schedule.remove(random.choice(rows))
        planted.append(("UNDERSTAFFED", f"{rid} is short by 1 person"))

for i, s in enumerate(schedule, 1):          # renumber so IDs have no gaps
    s["ScheduleID"] = f"S{i:03d}"

# ---------- 6. Training (2 courses per person = 80) ----------
training, tid = [], 1
for p in practitioners:
    for course in random.sample(COURSES, 2):
        due = date(2026, 9, 15) + timedelta(days=random.randint(0, 105))
        w = [70, 15, 15] if due < TODAY else [25, 30, 45]
        status = random.choices(["Completed", "In Progress", "Not Started"], weights=w)[0]
        comp = (min(due - timedelta(days=random.randint(0, 10)),
                    TODAY - timedelta(days=random.randint(0, 3)))
                if status == "Completed" else None)
        training.append({"TrainingID": f"T{tid:03d}", "PractID": p["PractID"],
                         "CourseName": course, "DueDate": due,
                         "Status": status, "CompletionDate": comp})
        tid += 1

# ---------- 7. Save ----------
out = Path(__file__).resolve().parent.parent / "data"
out.mkdir(exist_ok=True)
sheets = {"Practitioners": practitioners, "Requests": requests, "Schedule": schedule,
          "Leave": leave, "Training": training}
file = out / "ResourceOps_Tracker_filled.xlsx"
with pd.ExcelWriter(file, engine="openpyxl", date_format="DD-MMM-YYYY",
                    datetime_format="DD-MMM-YYYY") as w:
    for name, rows in sheets.items():
        pd.DataFrame(rows).to_excel(w, sheet_name=name, index=False)
    pd.DataFrame({"Dashboard": ["Built in Step 6"]}).to_excel(w, sheet_name="Dashboard", index=False)
    for ws in w.book.worksheets:
        for c in ws[1]:
            c.font = Font(name="Arial", bold=True, color="FFFFFF")
            c.fill = PatternFill("solid", fgColor="1F4E78")
        for col in ws.columns:
            ws.column_dimensions[col[0].column_letter].width = 20
        ws.freeze_panes = "A2"

pd.DataFrame(planted, columns=["ErrorType", "Detail"]).to_csv(
    out / "planted_errors_DO_NOT_OPEN.csv", index=False)

print("Done. Rows created:")
for name, rows in sheets.items():
    print(f"  {name:14s}{len(rows)}")
print(f"  Planted errors: {len(planted)} (saved in data/planted_errors_DO_NOT_OPEN.csv)")
