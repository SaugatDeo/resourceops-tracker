"""exceptions_report.py - Monday exceptions report for the ResourceOps Tracker.
Reads the workbook, re-runs every check in pandas (independent of Excel formulas)
and writes a one-file Excel report.

Usage (from the ResourceOps folder):
    python scripts/exceptions_report.py
    python scripts/exceptions_report.py "data/ResourceOps_Tracker_filled 2.xlsx"
    python scripts/exceptions_report.py --today 2026-10-07
"""
import argparse
from pathlib import Path

import numpy as np
import pandas as pd
from openpyxl.styles import Font, PatternFill

SLA_DAYS = 5

parser = argparse.ArgumentParser()
parser.add_argument("file", nargs="?", help="path to the tracker .xlsx")
parser.add_argument("--today", help="override today's date, YYYY-MM-DD")
args = parser.parse_args()

data_dir = Path(__file__).resolve().parent.parent / "data"
if args.file:
    src = Path(args.file)
else:
    files = [f for f in data_dir.glob("*.xlsx") if not f.name.startswith("Monday_Exceptions")]
    if not files:
        raise SystemExit("No .xlsx found in data/. Pass the file path as an argument.")
    src = max(files, key=lambda f: f.stat().st_mtime)

today = pd.Timestamp(args.today) if args.today else pd.Timestamp.today().normalize()
print(f"Reading: {src.name}   |   Report date: {today:%d-%b-%Y}")

def load(sheet):
    return pd.read_excel(src, sheet_name=sheet)

pract, req = load("Practitioners"), load("Requests")
sched, leave, train = load("Schedule"), load("Leave"), load("Training")
for df, cols in [(req, ["EventDate", "DateReceived", "DateClosed"]),
                 (sched, ["Date"]), (leave, ["LeaveStart", "LeaveEnd"]),
                 (train, ["DueDate", "CompletionDate"])]:
    for c in cols:
        df[c] = pd.to_datetime(df[c], errors="coerce")

names = pract.set_index("PractID")["Name"]

# 1. Double bookings: same person, same date, more than one schedule row
dup = sched[sched.duplicated(["PractID", "Date"], keep=False)].copy()
dup["Name"] = dup["PractID"].map(names)
double_bookings = dup.sort_values(["PractID", "Date"])[
    ["PractID", "Name", "Date", "RequestID", "ScheduleID"]]

# 2. Leave clashes: scheduled inside an approved leave window
m = sched.merge(leave, on="PractID")
clash = m[(m["LeaveStart"] <= m["Date"]) & (m["Date"] <= m["LeaveEnd"])].copy()
clash["Name"] = clash["PractID"].map(names)
leave_clashes = clash.drop_duplicates("ScheduleID")[
    ["PractID", "Name", "Date", "RequestID", "LeaveStart", "LeaveEnd", "LeaveType"]]

# 3. Understaffed requests (Open requests are pending, not errors)
assigned = sched.groupby("RequestID").size()
r = req[["RequestID", "Client", "Region", "Location", "EventDate",
         "PeopleNeeded", "Status"]].copy()          # ignore Excel helper columns
r["StaffAssigned"] = r["RequestID"].map(assigned).fillna(0).astype(int)
r["StaffGap"] = r["PeopleNeeded"] - r["StaffAssigned"]
understaffed = r[(r["Status"] != "Open") & (r["StaffGap"] > 0)][
    ["RequestID", "Client", "Region", "Location", "EventDate",
     "PeopleNeeded", "StaffAssigned", "StaffGap"]]

# 4. Overdue mandatory training
overdue = train[(train["Status"] != "Completed") & (train["DueDate"] < today)].copy()
overdue["Name"] = overdue["PractID"].map(names)
overdue["DaysOverdue"] = (today - overdue["DueDate"]).dt.days
overdue = overdue.sort_values("DaysOverdue", ascending=False)[
    ["PractID", "Name", "CourseName", "DueDate", "Status", "DaysOverdue"]]

# 5. SLA breaches on closed requests (working days from received to closed)
closed = req[(req["Status"] == "Closed") & req["DateClosed"].notna()].copy()
closed["WorkingDays"] = [
    int(np.busday_count(a.date(), b.date()))
    for a, b in zip(closed["DateReceived"], closed["DateClosed"])]
sla_breaches = closed[closed["WorkingDays"] > SLA_DAYS][
    ["RequestID", "Client", "DateReceived", "DateClosed", "WorkingDays"]]
met = len(closed) - len(sla_breaches)

summary = pd.DataFrame({
    "Check": ["Double-booked rows", "Leave clashes", "Understaffed requests",
              "Overdue training", f"SLA breaches (>{SLA_DAYS} working days)",
              "Closed requests within SLA"],
    "Result": [len(double_bookings), len(leave_clashes), len(understaffed),
               len(overdue), len(sla_breaches),
               f"{met} of {len(closed)} ({met / max(len(closed), 1):.1%})"],
})

out = data_dir / f"Monday_Exceptions_Report_{today:%Y-%m-%d}.xlsx"
sheets = {"Summary": summary, "DoubleBookings": double_bookings,
          "LeaveClashes": leave_clashes, "Understaffed": understaffed,
          "OverdueTraining": overdue, "SLABreaches": sla_breaches}
with pd.ExcelWriter(out, engine="openpyxl", date_format="DD-MMM-YYYY",
                    datetime_format="DD-MMM-YYYY") as w:
    for name, df in sheets.items():
        df.to_excel(w, sheet_name=name, index=False)
    for ws in w.book.worksheets:
        for c in ws[1]:
            c.font = Font(bold=True, color="FFFFFF")
            c.fill = PatternFill("solid", fgColor="1F4E78")
        for col in ws.columns:
            ws.column_dimensions[col[0].column_letter].width = 22
        ws.freeze_panes = "A2"

print("\nMONDAY EXCEPTIONS SUMMARY")
print(summary.to_string(index=False))
print(f"\nSaved: {out}")
